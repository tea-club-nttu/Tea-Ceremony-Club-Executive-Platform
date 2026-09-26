import re
import unittest
from io import BytesIO
from itertools import product
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from utils.achievement_report import build_report
from utils.application_form import build_application_form
from utils.teacher_comment import generate_application_progress_with_preview


class DocumentFieldsTest(unittest.TestCase):
    def test_revised_application_template(self):
        names = {role: f"測試{role}" for role in ("社長", "副社長", "總務", "攝錄", "點心", "文書", "美宣")}
        output = build_application_form(template_file=None, fields={
            "officer_names": names, "activity_date": "活動日期測試", "progress_date": "流程日期測試",
        })
        doc = Document(output)
        text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        for name in names.values():
            self.assertIn(name, text)
        self.assertNotIn("{{", text)
        self.assertIn("活動日期測試", text)
        self.assertEqual(text.count("流程日期測試"), 4)

    def test_report_academic_fields(self):
        output, _ = build_report(template_file=None, questionnaire_file=None,
                                 fields={"academic_year": "115", "semester": "一"}, images={})
        text = "\n".join(p.text for p in Document(output).paragraphs)
        self.assertIn("115學年度第一學期", text)

    def test_application_fields_preserve_format(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("{{社長姓名}} {{點心姓名}} {{活動流程日期}} {{點心}}")
        run.font.size = Pt(16)
        run.bold = True
        source = BytesIO()
        doc.save(source)
        source.seek(0)
        output = build_application_form(template_file=source, fields={
            "officer_names": {"社長": "甲", "點心": "乙"},
            "progress_date": "115 年 9 月 26 日", "snack_item": "餅乾",
        })
        paragraph = Document(output).paragraphs[0]
        self.assertEqual(paragraph.text, "甲 乙 115 年 9 月 26 日 餅乾")
        self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(paragraph.runs[0].font.size, Pt(16))
        self.assertTrue(paragraph.runs[0].bold)

    def test_progress_options_and_times(self):
        for ice, diy, health, video, song in product((False, True), repeat=5):
            result = generate_application_progress_with_preview(
                api_key=None, model="", activity_name="封箱茶會", tea_topic="烏龍茶", snack_item="餅乾",
                include_icebreaker=ice, include_snack_diy=diy, include_health_chat=health,
                include_review_video=video, include_club_song=song,
            )
            text = result["final_text"]
            self.assertEqual("影片" in text, video)
            self.assertEqual("社歌" in text, song)
            previous = 19 * 60 + 35
            for line in text.splitlines():
                sh, sm, eh, em = map(int, re.match(r"(\d+):(\d+)-(\d+):(\d+)", line).groups())
                start, end = sh * 60 + sm, eh * 60 + em
                self.assertEqual(start, previous)
                self.assertGreater(end, start)
                self.assertLessEqual(end, 20 * 60 + 50)
                previous = end

    def test_ai_cannot_add_unchecked_items(self):
        with patch("utils.teacher_comment.generate_ai_result", return_value={
            "text": "19:35-20:50 介紹茶、看回顧影片與唱社歌", "debug": {}, "provider": "test", "model": "test",
        }):
            result = generate_application_progress_with_preview(
                api_key="test", model="test", activity_name="封箱", tea_topic="茶", snack_item="",
            )
        self.assertNotIn("影片", result["final_text"])
        self.assertNotIn("社歌", result["final_text"])


if __name__ == "__main__":
    unittest.main()
