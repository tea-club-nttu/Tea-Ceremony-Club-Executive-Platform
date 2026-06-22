from __future__ import annotations

import base64
import json
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib import error, parse, request
from zoneinfo import ZoneInfo

from docx import Document

from utils.github_json_store import (
    GITHUB_BRANCH,
    github_contents_url,
    github_request,
    github_token,
    load_json,
    save_json,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = PROJECT_ROOT / "data"
TEMPLATE_DIR = DATA_DIR / "templates"
TEMPLATES_METADATA_PATH = DATA_DIR / "templates.json"
GITHUB_METADATA_PATH = "data/templates.json"
GITHUB_TEMPLATE_DIR = "data/templates"

TEMPLATE_CONFIGS = {
    "achievement_report": {
        "label": "成果書模板",
        "default_path": PROJECT_ROOT / "assets" / "成果書模板_已標記.docx",
        "default_name": "成果書模板_已標記.docx",
        "placeholders": [
            "{{填寫日期}}",
            "{{活動名稱}}",
            "{{活動地點}}",
            "{{活動日期}}",
            "{{活動負責人}}",
            "{{連絡電話}}",
            "{{參加人數}}",
            "{{活動內容概述}}",
            "{{問卷分析結果}}",
            "{{活動檢討}}",
            "{{照片1說明}}",
            "{{照片2說明}}",
            "{{照片3說明}}",
            "{{指導老師評語}}",
        ],
    },
    "application_form": {
        "label": "活動申請書模板",
        "default_path": PROJECT_ROOT / "assets" / "活動申請書模板.docx",
        "default_name": "活動申請書模板.docx",
        "placeholders": [
            "{{活動名稱}}",
            "{{活動日期}}",
            "{{活動負責人}}",
            "{{活動副負責人}}",
            "{{負責人電話}}",
            "{{活動宗旨}}",
            "{{活動進行}}",
            "{{點心}}",
        ],
    },
}


def local_template_path(key: str) -> Path:
    return TEMPLATE_DIR / f"{key}.docx"


def github_template_path(key: str) -> str:
    return f"{GITHUB_TEMPLATE_DIR}/{key}.docx"


def normalize_templates(data: object) -> dict[str, dict[str, object]]:
    if not isinstance(data, dict):
        return {}

    templates = {}
    for key, item in data.items():
        if key not in TEMPLATE_CONFIGS or not isinstance(item, dict):
            continue

        file_name = str(item.get("file_name", "")).strip()
        updated_at = str(item.get("updated_at", "")).strip()
        size = int(item.get("size", 0) or 0)
        path = str(item.get("path", "")).strip() or github_template_path(key)
        content_base64 = str(item.get("content_base64", "")).strip()

        if file_name and (path or content_base64):
            templates[key] = {
                "file_name": file_name,
                "updated_at": updated_at,
                "size": size,
                "path": path,
            }
            if content_base64:
                templates[key]["content_base64"] = content_base64

    return templates


def load_templates() -> dict[str, dict[str, object]]:
    if github_token():
        remote_templates = load_json(GITHUB_METADATA_PATH)
        if remote_templates is not None:
            return normalize_templates(remote_templates)

    if not TEMPLATES_METADATA_PATH.exists():
        return {}

    try:
        data = json.loads(TEMPLATES_METADATA_PATH.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}

    return normalize_templates(data)


def save_templates(templates: dict[str, dict[str, object]]) -> None:
    templates = normalize_templates(templates)

    if github_token() and save_json(GITHUB_METADATA_PATH, templates, "Update document template metadata"):
        return

    DATA_DIR.mkdir(exist_ok=True)
    TEMPLATES_METADATA_PATH.write_text(
        json.dumps(templates, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def validate_docx(data: bytes) -> None:
    Document(BytesIO(data))


def get_github_file_sha(path: str) -> str | None:
    token = github_token()
    if not token:
        return None

    url = f"{github_contents_url(path)}?ref={parse.quote(GITHUB_BRANCH)}"
    try:
        current_file = github_request("GET", url, token)
        return current_file.get("sha")
    except error.HTTPError as exc:
        if exc.code == 404:
            return None
        raise


def save_github_binary(path: str, data: bytes, message: str) -> bool:
    token = github_token()
    if not token:
        return False

    payload = {
        "message": message,
        "content": base64.b64encode(data).decode("ascii"),
        "branch": GITHUB_BRANCH,
    }

    try:
        sha = get_github_file_sha(path)
        if sha:
            payload["sha"] = sha
        github_request("PUT", github_contents_url(path), token, payload)
    except (OSError, error.HTTPError, json.JSONDecodeError):
        return False

    return True


def load_github_binary(path: str) -> bytes | None:
    token = github_token()
    if not token:
        return None

    url = f"{github_contents_url(path)}?ref={parse.quote(GITHUB_BRANCH)}"
    req = request.Request(
        url,
        method="GET",
        headers={
            "Accept": "application/vnd.github.raw",
            "Authorization": f"Bearer {token}",
            "X-GitHub-Api-Version": "2022-11-28",
        },
    )

    try:
        with request.urlopen(req, timeout=20) as response:
            data = response.read()
    except (OSError, error.HTTPError):
        return None

    try:
        metadata = json.loads(data.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return data

    content = metadata.get("content", "")
    if content:
        try:
            return base64.b64decode(content)
        except ValueError:
            return None

    return None


def save_uploaded_template(key: str, uploaded_file) -> None:
    if key not in TEMPLATE_CONFIGS:
        raise ValueError("未知的模板類型。")

    uploaded_file.seek(0)
    data = uploaded_file.getvalue()
    validate_docx(data)

    path = github_template_path(key)
    if github_token() and save_github_binary(path, data, f"Update {TEMPLATE_CONFIGS[key]['label']}"):
        pass
    else:
        TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
        local_template_path(key).write_bytes(data)

    templates = load_templates()
    templates[key] = {
        "file_name": uploaded_file.name,
        "updated_at": datetime.now(ZoneInfo("Asia/Taipei")).strftime("%Y-%m-%d %H:%M:%S"),
        "size": len(data),
        "path": path,
    }
    save_templates(templates)


def delete_template(key: str) -> None:
    templates = load_templates()
    if key in templates:
        del templates[key]
        save_templates(templates)

    try:
        local_template_path(key).unlink()
    except FileNotFoundError:
        pass


def get_template_record(key: str) -> dict[str, object] | None:
    return load_templates().get(key)


def get_stored_template_data(key: str, record: dict[str, object]) -> bytes | None:
    content_base64 = str(record.get("content_base64", "")).strip()
    if content_base64:
        try:
            return base64.b64decode(content_base64)
        except ValueError:
            return None

    if github_token():
        data = load_github_binary(str(record.get("path", github_template_path(key))))
        if data is not None:
            return data

    path = local_template_path(key)
    if path.exists():
        return path.read_bytes()

    return None


def get_template_bytes(key: str) -> tuple[bytes, str, bool]:
    config = TEMPLATE_CONFIGS[key]
    record = get_template_record(key)
    if record:
        data = get_stored_template_data(key, record)
        if data is not None:
            return data, str(record["file_name"]), True

    default_path = Path(config["default_path"])
    return default_path.read_bytes(), str(config["default_name"]), False


def get_template_source(key: str):
    record = get_template_record(key)
    if record:
        data = get_stored_template_data(key, record)
        if data is not None:
            return BytesIO(data)

    return str(TEMPLATE_CONFIGS[key]["default_path"])


def template_status_text(key: str) -> str:
    record = get_template_record(key)
    if not record:
        return "目前使用內建模板。"

    file_name = record.get("file_name", "")
    updated_at = record.get("updated_at", "")
    return f"目前使用自訂模板：{file_name}（{updated_at} 更新）"
